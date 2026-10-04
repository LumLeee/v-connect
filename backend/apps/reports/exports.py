from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .analytics import LABELS, STATUSES


def append_row(sheet, values):
    sheet.append([ILLEGAL_CHARACTERS_RE.sub('', value) if isinstance(value, str) else value for value in values])
    # User-controlled titles/names/filter text are strings, never spreadsheet formulas.
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str):
            cell.data_type = 's'


def style_table(sheet, widths):
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='185547')
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    sheet.row_dimensions[1].height = 42
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = '1:1'


def report_workbook(report):
    workbook = Workbook()
    summary = workbook.active
    summary.title = 'Tổng quan'
    append_row(summary, ['Chỉ tiêu', 'Giá trị'])
    metadata = [('Báo cáo', 'V-Connect — Báo cáo hoạt động'), ('Phạm vi', report['scope']),
                ('Lập lúc (giờ Việt Nam)', report['generated_at']), ('Số hoạt động', report['activity_count'])]
    filter_labels = {'search': 'Từ khóa', 'status': 'Trạng thái', 'date_from': 'Ngày bắt đầu từ',
                     'date_to': 'Ngày bắt đầu đến', 'organizer': 'ID Nhà tổ chức'}
    for key, value in report['filters'].items():
        metadata.append((filter_labels[key], STATUSES[value] if key == 'status' else value))
    for row in metadata:
        append_row(summary, row)
    for key, label in STATUSES.items():
        append_row(summary, [f'Hoạt động: {label}', report['activity_statuses'][key]])
    for key, label in LABELS.items():
        append_row(summary, [label, report['metrics'][key]])
    append_row(summary, ['Phản hồi hợp lệ', report['feedback']['count']])
    append_row(summary, ['Điểm trung bình', report['feedback']['average_rating']])
    for rating, count in report['feedback']['distribution'].items():
        append_row(summary, [f'Đánh giá {rating} sao', count])
    append_row(summary, ['Quy tắc', 'Mỗi người/hoạt động tính một đơn hiện tại; đăng ký lại không tăng số đơn. '
                'Đã tham gia chỉ tính Hoàn thành, được duyệt và có điểm danh. Phản hồi bị ẩn không tính. '
                'Lọc và nhóm tháng theo ngày bắt đầu hoạt động (giờ Việt Nam), không phải ngày gửi đơn.'])
    style_table(summary, [38, 105])
    summary.row_dimensions[summary.max_row].height = 60
    details = workbook.create_sheet('Hoạt động')
    append_row(details, ['ID', 'Tên hoạt động', 'Nhà tổ chức', 'Trạng thái', 'Bắt đầu (giờ VN)', 'Kết thúc (giờ VN)',
                        *LABELS.values(), 'Phản hồi hợp lệ', 'Điểm trung bình'])
    for row in report['rows']:
        append_row(details, [row['id'], row['title'], row['organizer_name'], STATUSES[row['status']],
                            datetime.fromisoformat(row['starts_at']).replace(tzinfo=None),
                            datetime.fromisoformat(row['ends_at']).replace(tzinfo=None),
                            *[row['metrics'][key] for key in LABELS], row['feedback']['count'], row['feedback']['average_rating']])
        for column in [5, 6]:
            details.cell(details.max_row, column).number_format = 'dd/mm/yyyy hh:mm'
        details.cell(details.max_row, 16).number_format = '0.00'
    style_table(details, [38, 45, 32, 18, 22, 22] + [18] * 10)
    monthly = workbook.create_sheet('Theo tháng')
    append_row(monthly, ['Tháng bắt đầu hoạt động', 'Số hoạt động', 'Tổng đơn đăng ký', 'Đã tham gia (hoàn thành)', 'Phản hồi hợp lệ'])
    for month in report['months']:
        append_row(monthly, [month[key] for key in ['month', 'activities', 'registered', 'attended_completed', 'feedback']])
    style_table(monthly, [28, 20, 24, 30, 24])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
