from datetime import datetime, timedelta, timezone as utc_timezone
from io import BytesIO
from unittest.mock import patch

from django.core.cache import cache
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient, APITestCase

from apps.accounts.models import User
from apps.activities.models import Activity
from apps.feedback import services as feedback_services
from apps.feedback.tests import fixtures
from apps.participations.models import Attendance, Participation
from .analytics import MAX_ACTIVITIES, build_report
from .services import metrics_for


class AnalyticsTests(APITestCase):
    url = '/api/v1/reports/analytics/'
    export = '/api/v1/reports/analytics/export/'

    def setUp(self):
        cache.clear()
        self.owner, self.one, self.two, self.admin, self.activity, self.attendances = fixtures()
        self.client = APIClient()
        self.client.force_login(self.owner)
        self.other = User.objects.create_user('analytics.other@example.invalid', full_name='Tổ chức khác', role='organizer')
        self.foreign = self.extra_activity(self.other)

    def extra_activity(self, owner=None, status='draft', title='Hoạt động khác'):
        now = timezone.now()
        return Activity.objects.create(organizer=owner or self.owner, title=title, description='Nội dung',
            address='Huế', capacity=20, starts_at=now - timedelta(days=1), ends_at=now - timedelta(hours=1),
            status=status, published_at=now - timedelta(days=2) if status != 'draft' else None)

    def add_attendance(self, activity, volunteer, status='approved'):
        entry = Participation.objects.create(activity=activity, volunteer=volunteer, status=status,
            registered_starts_at=activity.starts_at, registered_ends_at=activity.ends_at, registered_address=activity.address)
        return Attendance.objects.create(participation=entry, confirmed_by=activity.organizer)

    def test_totals_match_basic_reports_and_average_is_weighted(self):
        feedback_services.submit(self.activity.pk, self.one, {'rating': 5, 'content': 'Tốt'})
        feedback_services.submit(self.activity.pk, self.two, {'rating': 1, 'content': 'Cần cải thiện'})
        extra = self.extra_activity(status='completed')
        self.add_attendance(extra, self.one)
        feedback_services.submit(extra.pk, self.one, {'rating': 5, 'content': 'Hữu ích'})
        data = self.client.get(self.url).data
        basic = metrics_for(self.owner)
        self.assertEqual(data['metrics'], {key: value for key, value in basic.items() if key != 'feedback'})
        self.assertEqual(data['feedback']['average_rating'], 3.67)
        self.assertEqual(data['feedback']['count'], 3)
        self.assertEqual(data['feedback']['distribution']['5'], 2)
        self.assertEqual(sum(row['metrics']['registered'] for row in data['rows']), data['metrics']['registered'])
        self.assertEqual(sum(month['feedback'] for month in data['months']), 3)

    def test_hidden_feedback_and_cancelled_attendance_rules(self):
        feedback = feedback_services.submit(self.activity.pk, self.one, {'rating': 4, 'content': 'Hợp lệ'})
        feedback_services.hide(feedback.pk, self.admin, 'Ẩn theo yêu cầu')
        cancelled = self.extra_activity(status='cancelled')
        self.add_attendance(cancelled, self.one, status='cancelled')
        ongoing = self.extra_activity(status='published')
        self.add_attendance(ongoing, self.one)
        data = self.client.get(self.url).data
        self.assertEqual(data['feedback']['count'], 0)
        self.assertIsNone(data['feedback']['average_rating'])
        self.assertEqual(data['metrics']['attended_completed'], 2)
        self.assertEqual(data['metrics']['attended_cancelled'], 1)
        self.assertEqual(data['metrics']['attended_ongoing'], 1)

    def test_scope_and_export_permissions(self):
        for url in [self.url, self.export]:
            response = self.client.get(url, {'organizer': str(self.other.pk)})
            self.assertEqual(response.status_code, 403)
            self.client.force_login(self.one)
            self.assertEqual(self.client.get(url).status_code, 403)
            self.client.logout()
            self.assertEqual(self.client.get(url).status_code, 401)
            self.client.force_login(self.owner)
        data = self.client.get(self.url).data
        self.assertNotIn(str(self.foreign.pk), [row['id'] for row in data['rows']])
        self.client.force_login(self.admin)
        data = self.client.get(self.url, {'organizer': str(self.other.pk)}).data
        self.assertEqual([row['id'] for row in data['rows']], [str(self.foreign.pk)])
        self.assertEqual(data['scope'], self.other.full_name)
        self.assertEqual(self.client.get(self.url).data['activity_count'], 2)

    def test_filters_and_vietnam_month_boundaries(self):
        start = datetime(2027, 1, 31, 17, tzinfo=utc_timezone.utc)
        Activity.objects.filter(pk=self.activity.pk).update(starts_at=start, ends_at=start + timedelta(hours=1))
        query = {'date_from': '2027-02-01', 'date_to': '2027-02-01', 'status': 'completed', 'search': self.activity.title}
        data = self.client.get(self.url, query).data
        self.assertEqual(data['activity_count'], 1)
        self.assertEqual(data['months'][0]['month'], '2027-02')
        self.assertEqual(self.client.get(self.url, {**query, 'date_to': '2027-01-31'}).status_code, 400)
        self.assertEqual(self.client.get(self.url, {'date_to': '2027-01-31'}).data['activity_count'], 0)
        self.assertEqual(self.client.get(self.url, {**query, 'status': 'draft'}).data['activity_count'], 0)
        for url in [self.url, self.export]:
            for invalid in [{'date_from': 'bad'}, {'date_from': '0001-01-01'}, {'status': 'bad'}, {'search': 'x' * 201}]:
                self.assertEqual(self.client.get(url, invalid).status_code, 400)
        self.client.force_login(self.admin)
        for value in ['not-uuid', str(self.one.pk)]:
            self.assertEqual(self.client.get(self.url, {'organizer': value}).status_code, 400)

    def test_empty_report_and_limit_never_silently_truncate(self):
        data = self.client.get(self.url, {'search': 'không tồn tại'}).data
        self.assertEqual(data['activity_count'], 0)
        self.assertEqual(data['rows'], [])
        self.assertEqual(data['months'], [])
        self.assertTrue(all(value == 0 for value in data['metrics'].values()))
        self.assertIsNone(data['feedback']['average_rating'])
        with patch('apps.reports.analytics.MAX_ACTIVITIES', 0):
            for url in [self.url, self.export]:
                self.assertEqual(self.client.get(url).status_code, 400)
        self.assertGreater(MAX_ACTIVITIES, 20)

    def test_excel_matches_report_all_rows_and_preserves_text_safely(self):
        Activity.objects.filter(pk=self.activity.pk).update(title='=HYPERLINK("https://example.invalid","Không chạy")')
        for n in range(21):
            self.extra_activity(title=f'Dòng {n}')
        report = build_report(self.owner, {})
        response = self.client.get(self.export, {'page': 2, 'page_size': 1})
        self.assertEqual(response.status_code, 200)
        self.assertIn('no-store', response['Cache-Control'])
        self.assertIn('.xlsx', response['Content-Disposition'])
        workbook = load_workbook(BytesIO(response.content))
        self.assertEqual(workbook.sheetnames, ['Tổng quan', 'Hoạt động', 'Theo tháng'])
        sheet = workbook['Hoạt động']
        self.assertEqual(sheet.max_row, report['activity_count'] + 1)
        self.assertEqual(sum(row[6].value for row in list(sheet.rows)[1:]), report['metrics']['registered'])
        unsafe = [cell for cell in sheet['B'] if str(cell.value).startswith('=')]
        self.assertEqual(len(unsafe), 1)
        self.assertEqual(unsafe[0].data_type, 's')
        self.assertFalse(any(cell.data_type == 'f' for tab in workbook for row in tab for cell in row))
        self.assertEqual(sheet['E2'].number_format, 'dd/mm/yyyy hh:mm')
        self.assertEqual(sheet.freeze_panes, 'A2')
        workbook.close()

    def test_excel_obeys_filters_and_excludes_foreign_activity(self):
        response = self.client.get(self.export, {'status': 'draft'})
        workbook = load_workbook(BytesIO(response.content))
        self.assertEqual(workbook['Hoạt động'].max_row, 1)
        workbook.close()
        self.client.force_login(self.admin)
        response = self.client.get(self.export, {'organizer': str(self.other.pk)})
        workbook = load_workbook(BytesIO(response.content))
        self.assertEqual(workbook['Hoạt động'].max_row, 2)
        self.assertEqual(workbook['Hoạt động']['A2'].value, str(self.foreign.pk))
        workbook.close()

    def test_organizer_options_are_admin_only_paginated_without_contact_details(self):
        url = '/api/v1/reports/organizers/'
        self.assertEqual(self.client.get(url).status_code, 403)
        self.client.force_login(self.admin)
        result = self.client.get(url, {'page_size': 1})
        self.assertEqual(result.data['count'], 2)
        self.assertIsNotNone(result.data['next'])
        self.assertEqual(set(result.data['results'][0]), {'id', 'name', 'is_active'})
        self.assertIn('no-store', result['Cache-Control'])
